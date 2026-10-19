"""Builds the Seller Margin Toolkit workbook (works in Excel and Google Sheets)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

HDR = PatternFill("solid", fgColor="0A6B4F")
INP = PatternFill("solid", fgColor="FFF8E1")
WHITE = Font(bold=True, color="FFFFFF")
MONEY = '"$"#,##0.00'

wb = Workbook()

def header(ws, row, labels, widths):
    for i, (l, w) in enumerate(zip(labels, widths), 1):
        c = ws.cell(row=row, column=i, value=l)
        c.fill, c.font = HDR, WHITE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)

# ---- Start here
ws = wb.active; ws.title = "Start Here"
ws["A1"] = "Seller Margin Toolkit"; ws["A1"].font = Font(bold=True, size=18)
for i, t in enumerate([
    "Yellow cells are inputs. Everything else calculates itself.",
    "1. Settings: enter your marketplace fee rates (check each marketplace's official fee page; defaults are US estimates).",
    "2. Inventory: list your products and unit costs.",
    "3. Sales Log: add one row per sale. Fees and profit fill in automatically.",
    "4. Ad Spend: log ad spend to see return on ad spend (ROAS).",
    "5. Dashboard: monthly profit, margin and a quarterly tax set-aside estimate.",
    "Not financial or tax advice. Tax rate is your own assumption; ask an accountant.",
], 3):
    ws.cell(row=i, column=1, value=t)
ws.column_dimensions["A"].width = 110

# ---- Settings
st = wb.create_sheet("Settings")
header(st, 1, ["Marketplace", "Listing fee ($)", "Fee % of item+shipping", "Payment % ", "Payment fixed ($)", "Per-order fee ($)"],
       [24, 14, 22, 12, 16, 16])
rows = [("Etsy", 0.20, 6.5, 3.0, 0.25, 0), ("eBay", 0, 13.6, 0, 0, 0.40),
        ("Amazon", 0, 15, 0, 0, 0), ("Other", 0, 0, 0, 0, 0)]
for r, row in enumerate(rows, 2):
    for c, v in enumerate(row, 1):
        cell = st.cell(row=r, column=c, value=v)
        if c > 1: cell.fill = INP
st["A8"] = "Tax set-aside % of profit"; st["B8"] = 25; st["B8"].fill = INP
st["A9"] = "Fee defaults are estimates (2026-10-19). Verify before relying on them."

# ---- Inventory
inv = wb.create_sheet("Inventory")
header(inv, 1, ["SKU / Product", "Unit cost ($)", "Qty bought", "Qty sold", "In stock", "Reorder at", "Status"],
       [28, 14, 12, 10, 10, 11, 14])
for r in range(2, 202):
    inv[f"D{r}"] = f'=SUMIF(\'Sales Log\'!$B:$B,A{r},\'Sales Log\'!$D:$D)'
    inv[f"E{r}"] = f'=IF(A{r}="","",C{r}-D{r})'
    inv[f"G{r}"] = f'=IF(A{r}="","",IF(E{r}<=F{r},"REORDER","OK"))'
    for c in "ABCF": inv[f"{c}{r}"].fill = INP
    inv[f"B{r}"].number_format = MONEY
inv.conditional_formatting.add("G2:G201", CellIsRule(operator="equal", formula=['"REORDER"'],
                               fill=PatternFill("solid", bgColor="F8D7DA")))

# ---- Sales Log
sl = wb.create_sheet("Sales Log")
header(sl, 1, ["Date", "Product (match Inventory)", "Marketplace", "Qty", "Item price ($)", "Shipping charged ($)",
               "Your shipping cost ($)", "Item cost ($)", "Fees ($)", "Net profit ($)", "Margin"],
       [12, 28, 14, 7, 14, 16, 16, 13, 12, 14, 10])
dv = DataValidation(type="list", formula1="=Settings!$A$2:$A$5", allow_blank=True)
sl.add_data_validation(dv); dv.add("C2:C1001")
for r in range(2, 1002):
    sl[f"H{r}"] = f'=IFERROR(VLOOKUP(B{r},Inventory!$A:$B,2,FALSE)*D{r},0)'
    sl[f"I{r}"] = (f'=IF(C{r}="",0,IFERROR(VLOOKUP(C{r},Settings!$A$2:$F$5,2,FALSE)'
                   f'+(E{r}*D{r}+F{r})*VLOOKUP(C{r},Settings!$A$2:$F$5,3,FALSE)/100'
                   f'+(E{r}*D{r}+F{r})*VLOOKUP(C{r},Settings!$A$2:$F$5,4,FALSE)/100'
                   f'+VLOOKUP(C{r},Settings!$A$2:$F$5,5,FALSE)+VLOOKUP(C{r},Settings!$A$2:$F$5,6,FALSE),0))')
    sl[f"J{r}"] = f'=IF(B{r}="","",E{r}*D{r}+F{r}-I{r}-G{r}-H{r})'
    sl[f"K{r}"] = f'=IF(OR(B{r}="",E{r}*D{r}+F{r}=0),"",J{r}/(E{r}*D{r}+F{r}))'
    sl[f"A{r}"].number_format = "yyyy-mm-dd"
    for c in "EFGHIJ": sl[f"{c}{r}"].number_format = MONEY
    sl[f"K{r}"].number_format = "0.0%"
    for c in "ABCDEFG": sl[f"{c}{r}"].fill = INP

# ---- Ad Spend
ad = wb.create_sheet("Ad Spend")
header(ad, 1, ["Month (yyyy-mm)", "Marketplace", "Ad spend ($)", "Ad revenue ($)", "ROAS", "Verdict"], [16, 14, 14, 15, 10, 16])
for r in range(2, 62):
    ad[f"E{r}"] = f'=IF(C{r}>0,D{r}/C{r},"")'
    ad[f"F{r}"] = f'=IF(E{r}="","",IF(E{r}>=3,"Scale",IF(E{r}>=1.5,"Watch","Losing money")))'
    for c in "ABCD": ad[f"{c}{r}"].fill = INP
    for c in "CD": ad[f"{c}{r}"].number_format = MONEY
    ad[f"E{r}"].number_format = "0.00"

# ---- Dashboard
db = wb.create_sheet("Dashboard", 1)
db["A1"] = "Dashboard"; db["A1"].font = Font(bold=True, size=16)
db.column_dimensions["A"].width = 34; db.column_dimensions["B"].width = 18
items = [
    ("Total revenue", "=SUMPRODUCT('Sales Log'!E2:E1001,'Sales Log'!D2:D1001)+SUM('Sales Log'!F2:F1001)"),
    ("Total fees", "=SUM('Sales Log'!I2:I1001)"),
    ("Total net profit", "=SUM('Sales Log'!J2:J1001)"),
    ("Total ad spend", "=SUM('Ad Spend'!C2:C61)"),
    ("Profit after ads", "=B5-B6"),
    ("Overall margin", "=IF(B3=0,0,B7/B3)"),
    ("Tax set-aside (see Settings)", "=MAX(0,B7)*Settings!B8/100"),
    ("Per-quarter set-aside (total ÷ 4, use with a full year of data)", "=B9/4"),
]
for i, (l, f) in enumerate(items, 3):
    db[f"A{i}"] = l; db[f"B{i}"] = f
    db[f"B{i}"].number_format = "0.0%" if "margin" in l.lower() else MONEY
db["A12"] = "Items to reorder"; db["B12"] = '=COUNTIF(Inventory!G2:G201,"REORDER")'

wb.save("Seller-Margin-Toolkit.xlsx")
print("saved")
